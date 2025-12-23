"""
급여 명세서 생성기
"""
from datetime import datetime
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from employee import Employee


class PayrollStatementGenerator:
    """급여 명세서 생성"""

    def __init__(self, year: int, month: int):
        self.year = year
        self.month = month

    def generate(self, employees: List[Employee], output_path: str = "급여명세서.xlsx"):
        """급여 명세서 엑셀 파일 생성"""
        wb = Workbook()
        
        # 각 직원마다 별도 시트 생성
        for idx, employee in enumerate(employees):
            if idx == 0:
                ws = wb.active
                ws.title = employee.name[:30]  # 시트명 길이 제한
            else:
                ws = wb.create_sheet(title=employee.name[:30])
            
            self._create_statement_sheet(ws, employee)
        
        wb.save(output_path)
        return output_path

    def _create_statement_sheet(self, ws, employee: Employee):
        """개별 급여 명세서 시트 작성"""
        # 스타일 정의
        title_font = Font(size=16, bold=True)
        header_font = Font(size=11, bold=True)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
        
        # 제목
        ws.merge_cells('A1:F1')
        ws['A1'] = f'{self.year}년 {self.month}월 급여 명세서'
        ws['A1'].font = title_font
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        
        # 직원 정보
        ws['A3'] = '직원번호'
        ws['B3'] = employee.employee_id
        ws['D3'] = '성명'
        ws['E3'] = employee.name
        
        ws['A4'] = '부서'
        ws['B4'] = employee.department
        ws['D4'] = '직급'
        ws['E4'] = employee.position
        
        # 지급 항목 헤더
        ws['A6'] = '지급 항목'
        ws['B6'] = '금액'
        ws['A6'].font = header_font
        ws['B6'].font = header_font
        ws['A6'].fill = header_fill
        ws['B6'].fill = header_fill
        
        # 지급 항목
        row = 7
        ws[f'A{row}'] = '기본급'
        ws[f'B{row}'] = employee.base_salary
        row += 1
        
        if employee.bonus > 0:
            ws[f'A{row}'] = '상여금'
            ws[f'B{row}'] = employee.bonus
            row += 1
        
        if employee.overtime_pay > 0:
            ws[f'A{row}'] = '초과근무수당'
            ws[f'B{row}'] = employee.overtime_pay
            row += 1
        
        if employee.meal_allowance > 0:
            ws[f'A{row}'] = '식대'
            ws[f'B{row}'] = employee.meal_allowance
            row += 1
        
        if employee.transportation_allowance > 0:
            ws[f'A{row}'] = '교통비'
            ws[f'B{row}'] = employee.transportation_allowance
            row += 1
        
        # 지급 합계
        ws[f'A{row}'] = '지급 합계'
        ws[f'B{row}'] = employee.calculate_gross_salary()
        ws[f'A{row}'].font = header_font
        ws[f'B{row}'].font = header_font
        row += 2
        
        # 공제 항목 헤더
        ws[f'A{row}'] = '공제 항목'
        ws[f'B{row}'] = '금액'
        ws[f'A{row}'].font = header_font
        ws[f'B{row}'].font = header_font
        ws[f'A{row}'].fill = header_fill
        ws[f'B{row}'].fill = header_fill
        row += 1
        
        # 공제 항목
        if employee.national_pension > 0:
            ws[f'A{row}'] = '국민연금'
            ws[f'B{row}'] = employee.national_pension
            row += 1
        
        if employee.health_insurance > 0:
            ws[f'A{row}'] = '건강보험'
            ws[f'B{row}'] = employee.health_insurance
            row += 1
        
        if employee.employment_insurance > 0:
            ws[f'A{row}'] = '고용보험'
            ws[f'B{row}'] = employee.employment_insurance
            row += 1
        
        if employee.income_tax > 0:
            ws[f'A{row}'] = '소득세'
            ws[f'B{row}'] = employee.income_tax
            row += 1
        
        if employee.resident_tax > 0:
            ws[f'A{row}'] = '주민세'
            ws[f'B{row}'] = employee.resident_tax
            row += 1
        
        # 공제 합계
        ws[f'A{row}'] = '공제 합계'
        ws[f'B{row}'] = employee.calculate_total_deductions()
        ws[f'A{row}'].font = header_font
        ws[f'B{row}'].font = header_font
        row += 2
        
        # 실수령액
        ws[f'A{row}'] = '실수령액'
        ws[f'B{row}'] = employee.calculate_net_salary()
        ws[f'A{row}'].font = Font(size=12, bold=True)
        ws[f'B{row}'].font = Font(size=12, bold=True)
        ws[f'A{row}'].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        ws[f'B{row}'].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        
        # 열 너비 조정
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 15
        
        # 금액 셀 포맷 (숫자 형식)
        for row_idx in range(7, row + 1):
            if ws[f'B{row_idx}'].value and isinstance(ws[f'B{row_idx}'].value, int):
                ws[f'B{row_idx}'].number_format = '#,##0'
