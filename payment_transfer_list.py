"""
급여이체 리스트 생성기
"""
from datetime import datetime
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from employee import Employee


class PaymentTransferListGenerator:
    """급여이체 리스트 생성"""

    def __init__(self, year: int, month: int):
        self.year = year
        self.month = month

    def generate(self, employees: List[Employee], output_path: str = "급여이체리스트.xlsx"):
        """급여이체 리스트 엑셀 파일 생성"""
        wb = Workbook()
        ws = wb.active
        ws.title = "급여이체리스트"
        
        self._create_transfer_list(ws, employees)
        
        wb.save(output_path)
        return output_path

    def _create_transfer_list(self, ws, employees: List[Employee]):
        """급여이체 리스트 작성"""
        # 스타일 정의
        title_font = Font(size=14, bold=True)
        header_font = Font(size=11, bold=True)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font_white = Font(size=11, bold=True, color="FFFFFF")
        
        # 제목
        ws.merge_cells('A1:H1')
        ws['A1'] = f'{self.year}년 {self.month}월 급여이체 리스트'
        ws['A1'].font = title_font
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 25
        
        # 헤더
        headers = ['번호', '직원번호', '성명', '부서', '직급', '은행명', '계좌번호', '이체금액']
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=col_idx)
            cell.value = header
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
        
        # 데이터 입력
        total_amount = 0
        for idx, employee in enumerate(employees, start=1):
            row = 3 + idx
            net_salary = employee.calculate_net_salary()
            total_amount += net_salary
            
            ws.cell(row=row, column=1, value=idx)
            ws.cell(row=row, column=2, value=employee.employee_id)
            ws.cell(row=row, column=3, value=employee.name)
            ws.cell(row=row, column=4, value=employee.department)
            ws.cell(row=row, column=5, value=employee.position)
            ws.cell(row=row, column=6, value=employee.bank_name)
            ws.cell(row=row, column=7, value=employee.account_number)
            ws.cell(row=row, column=8, value=net_salary)
            
            # 금액 셀 포맷
            ws.cell(row=row, column=8).number_format = '#,##0'
            
            # 테두리 적용
            for col in range(1, 9):
                ws.cell(row=row, column=col).border = border
                ws.cell(row=row, column=col).alignment = Alignment(horizontal='center', vertical='center')
        
        # 합계 행
        summary_row = 3 + len(employees) + 1
        ws.merge_cells(f'A{summary_row}:G{summary_row}')
        ws[f'A{summary_row}'] = '합계'
        ws[f'A{summary_row}'].font = Font(size=11, bold=True)
        ws[f'A{summary_row}'].alignment = Alignment(horizontal='center', vertical='center')
        ws[f'A{summary_row}'].fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
        ws[f'H{summary_row}'] = total_amount
        ws[f'H{summary_row}'].font = Font(size=11, bold=True)
        ws[f'H{summary_row}'].number_format = '#,##0'
        ws[f'H{summary_row}'].fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
        
        # 테두리 적용
        for col in range(1, 9):
            ws.cell(row=summary_row, column=col).border = border
        
        # 열 너비 조정
        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 12
        ws.column_dimensions['F'].width = 12
        ws.column_dimensions['G'].width = 18
        ws.column_dimensions['H'].width = 15
